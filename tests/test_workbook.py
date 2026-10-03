"""Workbook regressions; all builds stay in Hermes scratch, never outputs/."""
import contextlib
import csv
import importlib.util
import shutil
import subprocess
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRATCH = Path.home() / "AppData/Local/hermes/cache/scratch"
SCRATCH.mkdir(parents=True, exist_ok=True)
spec = importlib.util.spec_from_file_location("build_workbook", ROOT / "src/build_workbook.py")
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=SCRATCH)
        self.addCleanup(self.temp.cleanup)
        self.out = Path(self.temp.name) / "check.xlsx"

    def build(self):
        with patch.object(builder, "OUT", self.out), contextlib.redirect_stdout(io.StringIO()):
            builder.main()
        return load_workbook(self.out)

    def test_libreoffice_calculates_default_quick_check_and_bridge(self):
        engine = shutil.which("soffice")
        if not engine:
            self.skipTest("LibreOffice not installed; openpyxl does not evaluate formulas")
        source = self.build()
        variants = [("recent", "2025 Q1", 0.9), ("early", "2015 Q1", ""),
                    ("invalid", "not a quarter", "Invalid quarter")]
        for name, selection, _ in variants:
            source["Quick check"]["B4"] = selection
            source.save(self.out.parent / f"{name}.xlsx")
        converted = self.out.parent / "calculated"
        profile = (self.out.parent / "lo-profile").as_uri()
        subprocess.run([engine, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "xlsx",
                        "--outdir", str(converted), str(self.out),
                        *[str(self.out.parent / f"{name}.xlsx") for name, _, _ in variants]],
                       check=True, capture_output=True, timeout=90)
        for name, _, expected in variants:
            evaluated = load_workbook(converted / f"{name}.xlsx", data_only=True)
            result = evaluated["Quick check"]["B12"].value
            if isinstance(expected, float):
                self.assertAlmostEqual(result, expected)
            elif expected == "":
                self.assertIn(result, (None, ""))
            else:
                self.assertEqual(result, expected)
        wb = load_workbook(converted / self.out.name, data_only=True)
        self.assertAlmostEqual(wb["Quick check"]["B12"].value, -0.1)
        self.assertAlmostEqual(wb["Quick check"]["B11"].value, 543.4)
        ws = wb["Quarterly"]
        values = {r[0]: r for r in ws.iter_rows(values_only=True) if r[0] != "Quarter"}
        self.assertAlmostEqual(values["2025 Q1"][9], 0.9)
        self.assertIn(values["2015 Q1"][9], (None, ""))
        with builder.BRIDGE_CSV.open(newline="", encoding="utf-8") as f:
            records = list(csv.DictReader(f))
        for i, record in enumerate(records, 3):
            for col, key in ((6, "delta_write_offs_sgd_m"), (7, "volume_sgd_m"), (8, "rate_sgd_m")):
                self.assertAlmostEqual(wb["Annual bridge"].cell(i, col).value, float(record[key]), delta=0.011)
            self.assertEqual(wb["Annual bridge"].cell(i, 9).value, 0)

    def test_pending_analysis_output_builds_workbook_with_blank_published_rate(self):
        import sys
        sys.path.insert(0, str(ROOT))
        from src import analysis
        annual = self.out.parent / 'annual.csv'
        with analysis.RAWA.open(newline='', encoding='utf-8') as file:
            rows = list(csv.reader(file))
        index = rows[0].index('2025')
        for row in rows:
            del row[index]
        with annual.open('w', newline='', encoding='utf-8') as file:
            csv.writer(file).writerows(rows)
        with patch.object(analysis, 'OUT', self.out.parent), patch.object(analysis, 'RAWA', annual), contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(SystemExit) as result:
                analysis.main()
        self.assertEqual(result.exception.code, 0)
        with patch.object(builder, 'BRIDGE_CSV', self.out.parent / 'yearly_bridge.csv'):
            wb = self.build()
        ws = wb['Annual bridge']
        latest = next(row for row in ws.iter_rows() if row[0].value == 2025)
        self.assertIsNone(latest[4].value)
        self.assertTrue(latest[5].value.startswith('='))

    def test_navigation_wrapping_and_print_layout(self):
        wb = self.build()
        for name in ("Quarterly", "Annual bridge"):
            ws = wb[name]
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertTrue(ws.auto_filter.ref)
            self.assertTrue(ws.tables)
            self.assertEqual(ws.print_title_rows, "$1:$1")
            self.assertTrue(ws["D1"].alignment.wrap_text)
        for ws in wb:
            self.assertTrue(ws.print_area)
            self.assertEqual(ws.page_setup.fitToWidth, 1)
        ws = wb["Quick check"]
        self.assertTrue(ws["A12"].alignment.wrap_text)
        self.assertTrue(ws["A13"].alignment.wrap_text)
        self.assertGreater(ws.row_dimensions[13].height, 30)

    def test_selector_rejects_invalid_entries_with_safe_lookup_fallback(self):
        wb = self.build()
        ws = wb["Quick check"]
        dv = list(ws.data_validations.dataValidation)[0]
        self.assertTrue(dv.showErrorMessage)
        self.assertEqual(dv.errorStyle, "stop")
        for row in range(7, 13):
            self.assertIn('"Invalid quarter"', ws.cell(row, 2).value)

    def test_annual_bridge_includes_baseline_and_all_csv_comparisons(self):
        wb = self.build()
        ws = wb["Annual bridge"]
        self.assertEqual(ws["A2"].value, 2015)
        self.assertIsNone(ws["F2"].value)
        with builder.BRIDGE_CSV.open(newline="", encoding="utf-8") as f:
            records = list(csv.DictReader(f))
        for i, record in enumerate(records, 3):
            self.assertEqual(ws.cell(i, 1).value, int(record["year"]))
            self.assertEqual(ws.cell(i, 6).value, f"=B{i}-B{i-1}")
            self.assertTrue(ws.cell(i, 9).value.startswith("=IF("))

    def test_failed_readback_preserves_existing_output(self):
        self.out.write_bytes(b"previous validated workbook")
        with patch.object(builder, "OUT", self.out), patch.object(builder, "load_workbook", side_effect=OSError("readback failed")), contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaisesRegex(OSError, "readback failed"):
                builder.main()
        self.assertEqual(self.out.read_bytes(), b"previous validated workbook")
        self.assertEqual(list(self.out.parent.iterdir()), [self.out])

    def test_normalizer_preserves_namespace_bindings(self):
        import zipfile
        import xml.etree.ElementTree as ET
        # lxml puts xsi locally on modified, not necessarily on the root.
        for prefix in ("dcterms", "ns0"):
            xml = (f'<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
                   f'xmlns:{prefix}="http://purl.org/dc/terms/">'
                   f'<{prefix}:modified xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
                   f'xsi:type="{prefix}:W3CDTF">2020-01-01T00:00:00Z</{prefix}:modified></cp:coreProperties>')
            with zipfile.ZipFile(self.out, "w") as z:
                z.writestr("docProps/core.xml", xml)
            builder.normalize_xlsx(self.out)
            with zipfile.ZipFile(self.out) as z:
                root = ET.fromstring(z.read("docProps/core.xml"))
            modified = root.find("{http://purl.org/dc/terms/}modified")
            self.assertEqual(modified.text, "2026-10-04T00:00:00Z")

    def test_rate_change_subtracts_current_and_prior_inputs(self):
        wb = self.build()
        ws = wb["Quarterly"]
        rows = {r[0].value: r for r in list(ws.rows)[1:]}
        for quarter, prior, expected in [("2026 Q2", "2025 Q2", -0.1),
                                          ("2025 Q1", "2024 Q1", 0.9)]:
            row = rows[quarter]
            self.assertAlmostEqual(row[6].value - rows[prior][6].value, expected)
            self.assertIn(f'$G{row[0].row}-SUMIFS(', row[9].value)
        self.assertEqual(wb["Quick check"]["B4"].value, "2026 Q2")
        self.assertAlmostEqual(rows["2026 Q2"][6].value, 5.9)
        self.assertAlmostEqual(rows["2025 Q2"][6].value, 6.0)
        self.assertAlmostEqual(sum(rows[q][5].value for q in ("2025 Q3", "2025 Q4", "2026 Q1", "2026 Q2")), 543.4)
        self.assertTrue(rows["2015 Q1"][9].value.endswith(',"")'))
        self.assertIn('>=5', rows["2015 Q1"][9].value)


if __name__ == "__main__":
    unittest.main()
