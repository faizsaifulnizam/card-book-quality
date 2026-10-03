"""Build a complete disposable run, validate, then publish generated directories.

Stage/validation failures leave published artifacts untouched. Publication rolls
back on ordinary exceptions, but directory swaps are NOT multi-file crash-atomic
and require a single writer (no concurrent pipeline/refresh runs). Power loss or
process kill during publication may require recovery from .publish-backup-*.
Raw acquisition is separate: run download.py --force explicitly to refresh.
"""
import csv
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from download import publish_paths

ROOT = Path(__file__).resolve().parent.parent
STAGES = ['download.py', 'build_dataset.py', 'audit.py', 'analysis.py', 'figures.py', 'build_workbook.py']
GENERATED = ['data/processed', 'outputs', 'reports/figures', 'docs/img']


def validate_outputs(root):
    """Read the completed artifacts; this is structural, not Excel recalculation."""
    import duckdb
    from openpyxl import load_workbook

    con = duckdb.connect()
    parquet = (root / 'data/processed/quarterly.parquet').as_posix().replace("'", "''")
    con.execute(f"CREATE VIEW quarterly AS SELECT * FROM read_parquet('{parquet}')")
    checks = con.execute((root / 'sql/03_checks.sql').read_text(encoding='utf-8')).fetchall()
    con.close()
    if any(n for _, n in checks):
        raise ValueError(f'completed parquet failed checks: {checks}')
    for name, key in [('yearly_bridge.csv', 'year'), ('book_split.csv', 'year'), ('sensitivity.csv', 'variant')]:
        with (root / 'outputs' / name).open(newline='', encoding='utf-8') as f:
            rows = list(csv.DictReader(f))
        if not rows or key not in rows[0]:
            raise ValueError(f'empty/malformed output: {name}')
    for i, name in enumerate(['timeline', 'bridge', 'contributions', 'book'], 1):
        for theme in ['', '-dark']:
            filename = f'f{i}_{name}{theme}.png'
            image = (root / 'reports/figures' / filename).read_bytes()
            if len(image) < 5000 or not image.startswith(b'\x89PNG\r\n\x1a\n'):
                raise ValueError(f'invalid image: {filename}')
            if image != (root / 'docs/img' / filename).read_bytes():
                raise ValueError(f'Pages image differs: {filename}')
    wb = load_workbook(root / 'outputs/quick_check.xlsx')
    if not {'Quarterly', 'Quick check', 'Annual bridge'} <= set(wb.sheetnames):
        raise ValueError('workbook missing required sheets')
    wb.close()


def main():
    sources = ['credit-charge-cards-quarterly.csv', 'credit-charge-cards-annual.csv', 'pull_manifest.json']
    if not all((ROOT / 'data/raw' / name).is_file() for name in sources):
        raise SystemExit('run_all requires vendored raw pair and manifest; run download.py --force separately')
    with tempfile.TemporaryDirectory(prefix='.pipeline-stage-', dir=ROOT.parent) as temp:
        staged = Path(temp) / 'repo'
        shutil.copytree(ROOT, staged, ignore=shutil.ignore_patterns('.git', '.venv', '__pycache__', '.publish-backup-*'))
        # Remove old generated files so omitted products cannot masquerade as success.
        # Keep non-generated Pages assets, which figures.py also copies when present.
        for relative in GENERATED[:-1]:
            shutil.rmtree(staged / relative, ignore_errors=True)
        for path in (staged / 'docs/img').glob('f[1-4]_*.png'):
            path.unlink()
        for stage in STAGES:
            print(f'\n===== {stage} =====', flush=True)
            result = subprocess.run([sys.executable, str(staged / 'src' / stage)], cwd=staged)
            if result.returncode:
                raise SystemExit(f'{stage} failed (exit {result.returncode}) — published outputs left untouched')
        validate_outputs(staged)
        publish_paths([(staged / relative, ROOT / relative) for relative in GENERATED])
    print('\nrun_all: complete validated output set published (exception rollback; not crash-atomic)')


if __name__ == '__main__':
    main()
