"""Build outputs/quick_check.xlsx — the formula-only Excel surface.

Sheets:
  Quick check    — pick a quarter from the dropdown → the block updates (XLOOKUP pulls +
                   a 4-quarter sum and a rate-change cell keyed to the Quarterly sheet's
                   Seq index — plain SUMIFS, no OFFSET, safe if rows are re-sorted)
  Quarterly      — the full quarterly table + Seq index (1 = oldest) + per-row windows:
                   last-4-quarter write-offs and rate change vs the same quarter a year earlier
  Annual bridge  — the rate-vs-volume split as LIVE formulas (SUMIFS / AVERAGEIFS from the
                   Quarterly sheet): write-offs, average rollover, recomputed rate, the two
                   bridge terms and a closure check; the published rate column is pasted
                   from the MAS annual file for comparison

Formula-only, no macros. Run (any cwd): python src/build_workbook.py
Receipts printed: row counts, the values a reviewer sees for the default selection, a
no-OFFSET scan across every formula, and a check that the 4-quarter formula's inputs
(the Seq window on the Quarterly sheet) sum to the value computed in Python.
"""
import csv
import xml.etree.ElementTree as ET
import sys
import zipfile
import tempfile
from datetime import datetime
from pathlib import Path

import duckdb
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
PARQUET = (ROOT / "data/processed/quarterly.parquet").as_posix()
BRIDGE_CSV = ROOT / "outputs/yearly_bridge.csv"
OUT = ROOT / "outputs/quick_check.xlsx"

INK = "14293D"
MUTED = "5C6B79"
SAND = "E7E3DC"
STAMP = datetime(2026, 10, 4, 0, 0, 0)

HEAD = Font(bold=True, color=INK)
NOTE = Font(size=9, color=MUTED)
FILL = PatternFill("solid", fgColor=SAND)


def label(qq):
    return f"{qq.year} Q{(qq.month - 1) // 3 + 1}"


def q(con, sql):
    return con.sql(sql).fetchall()


def normalize_xlsx(path):
    """Rewrite the saved workbook deterministically.

    openpyxl writes zip entries with the current clock time and re-stamps
    docProps/core.xml 'modified' at save (no config knob), so two builds of the
    same data differ in bytes. Rebuild the archive with a fixed entry order and
    timestamp: a reviewer re-running this script must see a byte-identical file.
    """
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in sorted(zin.namelist()):
            data = zin.read(name)
            if name == "docProps/core.xml":
                # Trusted XML generated locally by openpyxl, not uploaded XML.
                namespaces = {"cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                              "dc": "http://purl.org/dc/elements/1.1/",
                              "dcterms": "http://purl.org/dc/terms/",
                              "xsi": "http://www.w3.org/2001/XMLSchema-instance"}
                for prefix, uri in namespaces.items():
                    ET.register_namespace(prefix, uri)
                core = ET.fromstring(data)
                for tag in ("created", "modified"):
                    element = core.find(f'{{{namespaces["dcterms"]}}}{tag}')
                    if element is not None:
                        element.text = STAMP.isoformat() + "Z"
                        element.set(f'{{{namespaces["xsi"]}}}type', "dcterms:W3CDTF")
                data = ET.tostring(core, encoding="utf-8")
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, data)
    tmp.replace(path)


def main():
    # Same filesystem: validated replacement is atomic; failed builds are discarded.
    with tempfile.TemporaryDirectory(prefix=".workbook-", dir=OUT.parent) as tmp:
        candidate = Path(tmp) / OUT.name
        build(candidate)
        candidate.replace(OUT)
    print(f"wrote: {OUT.as_posix()}  ({OUT.stat().st_size} bytes)")


def build(out):
    if not Path(PARQUET).exists():
        sys.exit("data/processed/quarterly.parquet missing — run src/build_dataset.py first")
    con = duckdb.connect()
    rows = q(con, f"""SELECT quarter, billings_sgd_m, rollover_sgd_m, write_offs_sgd_m, charge_off_rate_pct
                      FROM read_parquet('{PARQUET}') ORDER BY quarter DESC""")
    quarters = [(label(r[0]), r[0].year, (r[0].month - 1) // 3 + 1,
                 float(r[1]), float(r[2]), float(r[3]), float(r[4])) for r in rows]
    n = len(quarters)
    last = n + 1                       # last data row on the Quarterly sheet (header = row 1)
    seq_of = {qq[0]: n - i for i, qq in enumerate(quarters)}  # 1 = oldest quarter

    wb = Workbook()
    # Deterministic output: fixed document timestamps (a rebuild must be byte-identical,
    # not "same bytes except the save time" — reviewers diff regenerated artifacts).
    wb.properties.created = STAMP
    wb.properties.modified = STAMP
    wb.properties.creator = "card-book-quality — src/build_workbook.py"
    wb.properties.lastModifiedBy = "card-book-quality — src/build_workbook.py"
    ws = wb.active
    ws.title = "Quick check"

    ws["A1"] = "Card book — quick check (MAS credit & charge cards, quarterly)"
    ws["A1"].font = Font(bold=True, size=13, color=INK)
    ws["A2"] = ("All issuers in Singapore, not one bank. Values S$M unless stated. "
                "Pick a quarter below — the block updates (formulas, no macros).")
    ws["A2"].font = NOTE

    ws["A4"] = "Quarter:"
    ws["A4"].font = HEAD
    ws["B4"] = quarters[0][0]  # default: latest
    ws["B4"].font = Font(bold=True)
    ws["B4"].fill = PatternFill("solid", fgColor="F3EFE7")

    dv = DataValidation(type="list", formula1=f"Quarterly!$A$2:$A${last}", allow_blank=False)
    dv.error = "Pick a quarter from the list"
    dv.errorStyle = "stop"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(ws["B4"])

    LBL = [
        ("Quarter", '=$B$4', None),
        ("Total billings (S$M, that quarter)", '=XLOOKUP($B$4,Quarterly!$A:$A,Quarterly!$D:$D)', "0.0"),
        ("Rollover balance (S$M, quarter-end)", '=XLOOKUP($B$4,Quarterly!$A:$A,Quarterly!$E:$E)', "0.0"),
        ("Write-offs (S$M, that quarter)", '=XLOOKUP($B$4,Quarterly!$A:$A,Quarterly!$F:$F)', "0.0"),
        ("Published charge-off rate (%, annualised)", '=XLOOKUP($B$4,Quarterly!$A:$A,Quarterly!$G:$G)', '0.0"%"'),
        ("Write-offs, last 4 quarters incl. selected (S$M)", '=XLOOKUP($B$4,Quarterly!$A:$A,Quarterly!$I:$I)', "0.0"),
        ("Rate change vs same quarter a year earlier (pt)",
         '=XLOOKUP($B$4,Quarterly!$A:$A,Quarterly!$J:$J)', "+0.0;-0.0"),
    ]
    r0 = 6
    for i, (name, formula, numfmt) in enumerate(LBL):
        r = r0 + i
        ws.cell(row=r, column=1, value=name).font = HEAD if i == 0 else Font(color=INK)
        if formula.startswith("=XLOOKUP("):
            formula = formula.replace("XLOOKUP(", "_xlfn.XLOOKUP(", 1)[:-1] + ',"Invalid quarter")'
        c = ws.cell(row=r, column=2, value=formula)
        c.font = Font(bold=(i == 0), color=INK)
        if numfmt:
            c.number_format = numfmt
    ws.cell(row=r0 + 7, column=1,
            value=("(blank 4-quarter sum = fewer than 4 quarters of history up to the selection; blank rate change = "
                   "no same quarter a year earlier — formula-blank, never #REF)")).font = NOTE

    notes = [
        "",
        "Refresh: re-run the pipeline, then re-run this builder —",
        "  python src/run_all.py             # everything in order (download → dataset → analysis → figures → workbook)",
        "  python src/build_workbook.py      # just this file",
        "",
        "The 4-quarter and rate-change cells are keyed to the Seq column (1 = oldest) on the Quarterly sheet —",
        "plain SUMIFS with a bounds guard: no OFFSET, and still correct if you re-sort the sheet.",
        "The Annual bridge sheet computes the split live with SUMIFS / AVERAGEIFS — it mirrors outputs/yearly_bridge.csv.",
        "Source: MAS credit & charge cards via data.gov.sg — Singapore Open Data Licence. See docs/data_audit.md.",
        "Charge-off rate is as published: bad debts written off ÷ average rollover balance, annualised.",
        "Formula-only, no macros. XLOOKUP needs Excel 2021+/365; older readers can swap it for INDEX/MATCH.",
        "Numbers reproduce outputs/yearly_bridge.csv and the README headline for the same pull.",
    ]
    nr = r0 + 9
    for i, t in enumerate(notes):
        ws.cell(row=nr + i, column=1, value=t).font = NOTE
        ws.cell(row=nr + i, column=1).alignment = Alignment(horizontal="left")

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 22

    # ---- Quarterly sheet: data + Seq index + per-row windows ----
    wq = wb.create_sheet("Quarterly")
    head = ["Quarter", "Year", "Qtr", "Total billings (S$M)", "Rollover balance (S$M)",
            "Write-offs (S$M)", "Published rate (%)", "Seq (1=oldest)",
            "Last 4 qtrs: write-offs (S$M)", "Rate chg vs yr earlier (pt)"]
    for j, h in enumerate(head, start=1):
        c = wq.cell(row=1, column=j, value=h)
        c.font = HEAD
        c.fill = FILL
    for i, row in enumerate(quarters, start=2):
        for j, v in enumerate(list(row) + [seq_of[row[0]]], start=1):
            wq.cell(row=i, column=j, value=v)
        wq.cell(row=i, column=9,
                value=f'=IF($H{i}>=4,SUMIFS($F$2:$F${last},$H$2:$H${last},">="&$H{i}-3,$H$2:$H${last},"<="&$H{i}),"")')
        wq.cell(row=i, column=10,
                value=f'=IF($H{i}>=5,$G{i}-SUMIFS($G$2:$G${last},$H$2:$H${last},$H{i}-4),"")')
    for col, w in zip("ABCDEFGHIJ", (10, 7, 5, 21, 22, 18, 18, 12, 26, 24)):
        wq.column_dimensions[col].width = w

    # ---- Annual bridge sheet: the split as live formulas ----
    wb2 = wb.create_sheet("Annual bridge")
    with BRIDGE_CSV.open(newline="", encoding="utf-8") as f:
        brows = list(csv.DictReader(f))
    baseline = int(brows[0]["previous_year"])
    assert sum(qq[1] == baseline for qq in quarters) == 4, "Incomplete annual baseline"
    brows.insert(0, {"year": baseline, "rate_pct_pub_annual": None})
    bhead = ["Year", "Write-offs (S$M)", "Avg rollover (S$M)", "Rate, recomputed (%)",
             "Rate, published (%)", "Δ write-offs (S$M)", "Volume effect (S$M)",
             "Rate effect (S$M)", "Closure (Δ−vol−rate, 0 = exact)"]
    for j, h in enumerate(bhead, start=1):
        c = wb2.cell(row=1, column=j, value=h)
        c.font = HEAD
        c.fill = FILL
    for i, row in enumerate(brows, start=2):
        wb2.cell(row=i, column=1, value=int(row["year"]))
        wb2.cell(row=i, column=2, value=f'=SUMIFS(Quarterly!$F$2:$F${last},Quarterly!$B$2:$B${last},$A{i})')
        wb2.cell(row=i, column=3, value=f'=AVERAGEIFS(Quarterly!$E$2:$E${last},Quarterly!$B$2:$B${last},$A{i})')
        wb2.cell(row=i, column=4, value=f"=B{i}/C{i}*100")
        if row["rate_pct_pub_annual"] not in (None, ""):
            wb2.cell(row=i, column=5, value=float(row["rate_pct_pub_annual"]))
        if i > 2:
            wb2.cell(row=i, column=6, value=f"=B{i}-B{i - 1}")
            wb2.cell(row=i, column=7, value=f"=(C{i}-C{i - 1})*(D{i - 1}+D{i})/200")
            wb2.cell(row=i, column=8, value=f"=(D{i}-D{i - 1})/100*(C{i - 1}+C{i})/2")
            wb2.cell(row=i, column=9, value=f'=IF(F{i}="","",ROUND(F{i}-G{i}-H{i},6))')
        for col, fmt in ((2, "0.0"), (3, "0.0"), (4, "0.00"), (5, "0.0"),
                         (6, "+0.0;-0.0"), (7, "+0.0;-0.0"), (8, "+0.0;-0.0"), (9, "0.000000")):
            wb2.cell(row=i, column=col).number_format = fmt
    for col, w in zip("ABCDEFGHI", (6, 15, 17, 18, 17, 16, 17, 16, 27)):
        wb2.column_dimensions[col].width = w
    note_row = len(brows) + 3
    wb2.cell(row=note_row, column=1,
             value="Live formulas from the Quarterly sheet (SUMIFS / AVERAGEIFS) — the same midpoint split as outputs/yearly_bridge.csv.").font = NOTE
    wb2.cell(row=note_row + 1, column=1,
             value="Column E (published rate) is pasted from the MAS annual file; volume = Δrollover × avg rate · rate = Δrate × avg rollover.").font = NOTE

    # Report layout: notes stay within the page; data tables keep their headers.
    for row in (1, 2, 13, *range(nr, nr + len(notes))):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 42 if row in (2, 13) else 30
    for row in range(6, 13):
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 30
    for sheet, ref, name in ((wq, f"A1:J{last}", "QuarterlyData"),
                              (wb2, f"A1:I{len(brows) + 1}", "AnnualBridgeData")):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = ref
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        sheet.print_title_rows = "1:1"
        sheet.row_dimensions[1].height = 42
        for cell in sheet[1]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in (note_row, note_row + 1):
        wb2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        wb2.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        wb2.row_dimensions[row].height = 30
    for sheet in wb:
        sheet.print_area = sheet.dimensions
        sheet.page_setup.orientation = "portrait" if sheet == ws else "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)
    normalize_xlsx(out)
    print(f"sheets: {wb.sheetnames} · quarterly rows: {len(quarters)} · bridge rows: {len(brows)}")

    # ------------------------------------------------------------------
    # Receipts: read the written file back and check what a reviewer gets.
    # openpyxl cannot evaluate Excel formulas — but it CAN verify every
    # formula's construction, and the numbers the formulas will compute are
    # reproduced in Python from the same rows they range over.
    # ------------------------------------------------------------------
    rb = load_workbook(out)

    offenders = [f"{sh.title}!{c.coordinate}"
                 for sh in rb.worksheets for row in sh.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith("=") and "OFFSET" in c.value.upper()]
    assert not offenders, f"OFFSET found in: {offenders}"
    print("no-OFFSET scan: PASS (0 formulas use OFFSET)")

    sel = rb["Quick check"]["B4"].value
    latest = quarters[0]
    seq_latest = seq_of[latest[0]]
    rolling4 = sum(x[5] for x in quarters[:4])
    f4_input = sum(x[5] for x in quarters if seq_latest - 3 <= seq_of[x[0]] <= seq_latest)
    prev_year_same = next(x for x in quarters if x[0] == f"{int(latest[1]) - 1} Q{latest[2]}")
    yoy = round(latest[6] - prev_year_same[6], 1)
    print("read-back OK — openpyxl sees:")
    print(f"  default selection: {sel}")
    print(f"  expected (what the formulas compute): billings {latest[3]:,.1f} · rollover {latest[4]:,.1f} · "
          f"write-offs {latest[5]:,.1f} · rate {latest[6]:.1f}% · 4q write-offs {rolling4:,.1f} · rate YoY {yoy:+.1f} pt")
    assert abs(f4_input - rolling4) < 1e-9
    print(f"  formula-input check: the SUMIFS row window (Seq {seq_latest - 3}–{seq_latest}) covers write-offs summing to "
          f"{f4_input:,.1f} S$M — equals the Python value (this pull's README number)")
    f11, f12 = rb["Quick check"]["B11"].value, rb["Quick check"]["B12"].value
    assert "Quarterly!$I:$I" in f11 and "Quarterly!$J:$J" in f12
    assert rb["Quick check"]["A12"].value == "Rate change vs same quarter a year earlier (pt)", "label row clobbered"
    assert rb["Quick check"]["A13"].value.startswith("(blank"), "caveat must sit on its own row"
    assert rb["Quick check"]["B4"].value == quarters[0][0]
    assert rb["Quarterly"]["A2"].value == quarters[0][0] and rb["Quarterly"][f"A{last}"].value == quarters[-1][0]
    col_a = [c.value for c in rb["Quarterly"]["A"] if c.value]
    assert col_a[0] == "Quarter" and len(col_a) == len(quarters) + 1
    seq_col = [rb["Quarterly"][f"H{r}"].value for r in range(2, last + 1)]
    assert seq_col[0] == n and seq_col[-1] == 1 and sorted(seq_col) == list(range(1, n + 1)), "Seq column broken"
    # Check the saved formulas, not merely a separately printed Python result.
    for r in range(2, last + 1):
        expected = f'=IF($H{r}>=5,$G{r}-SUMIFS($G$2:$G${last},$H$2:$H${last},$H{r}-4),"")'
        assert rb["Quarterly"].cell(r, 10).value == expected, f"Rate-change formula broken at J{r}"
    saved = {rb["Quarterly"].cell(r, 8).value: rb["Quarterly"].cell(r, 7).value
             for r in range(2, last + 1)}
    saved_yoy = round(saved[seq_latest] - saved[seq_latest - 4], 1)
    assert saved_yoy == yoy, "Saved rate-change inputs disagree with source"
    print(f"  rate-change formula/input check: {saved[seq_latest]:.1f} - {saved[seq_latest - 4]:.1f} = {saved_yoy:+.1f} pt (not engine evaluation)")
    print(f"  quarterly sheet: {n} rows · Seq 1..{n} (newest first) · per-row windows in cols I/J")
    print("RESULT: workbook written + read-back checks PASS")


if __name__ == "__main__":
    main()
